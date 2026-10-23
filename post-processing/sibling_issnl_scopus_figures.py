#!/usr/bin/env python3
"""Three groups of counts, written to numbers/sibling_issnl_scopus_figures.numbers.json (keys sisf_*) and
EVAL_OUTPUT_DIR/output_sibling_issnl_scopus_figures/sibling_issnl_scopus_figures.json.
- Sibling-inferred files (files without a citation log whose submission has another file with a log, as defined in
   stats_corpus.py): the entries promoted from `excluded_sibling_inferred_uncited` to confirmed-cited, with their files,
   submissions and strict / broad flags after the identity check, and the sibling files that hold no excluded entry.
- ISSN-L and cache state of the identity audit: the mismatch references with a "cache ISSN differs" vote versus only
   "null cache ISSN" votes; distinct ISSN pairs and their ISSN-L resolution; the references included again by the ISSN-L
   check; the references whose citing ISSN is only on cache rows created after the start of the assessment run; strict and
   broad counts with each group, and with both, included again.
   The tool's cache is a live database; the check asks whether an ISSN that appears only on rows created after the start of
   the assessment run (EVAL_ASSESSMENT_RUN_START, "YYYY-MM-DD HH:MM:SS", compared with the creation time of each cache row that
   carries the ISSN) changes the identity check. Needs EVAL_CACHE_DB_PATH_EARLIER (a copy of the cache taken before those rows
   were created) next to EVAL_CACHE_DB_PATH.
- Scopus source list (EVAL_SCOPUS_XLSX_PATH): rows, active and inactive
   titles and quality-flagged rows under the row rules of the tool's scopus.py; the active count equals
   source_updates.records_updated of source 'scopus' in the cache.
Reads EVAL_DB_PATH_BBL_TEX, the audit keys file, EVAL_ISSN_L_CHECK_JSON (issn_l_family_check.py output) and both caches. Stops
if a stored value does not reproduce."""
from __future__ import annotations

import json
import os
import sqlite3
from collections import defaultdict

import stats_common
from identity_vote_rules import build_any_row_has_issn_index, normalize_issn
from stats_common import required_path_env, submission_of

DB_PATH_BBL_TEX = required_path_env("EVAL_DB_PATH_BBL_TEX")
EVAL_OUTPUT_DIR = required_path_env("EVAL_OUTPUT_DIR")
AUDIT_KEYS_JSON = EVAL_OUTPUT_DIR / "output_identity_mismatch_audit" / "full_corpus_identity_mismatch_audit_keys.json"
# output of issn_l_family_check.py; it must be on the basis of the audit keys file
ISSN_L_CHECK_JSON = required_path_env("EVAL_ISSN_L_CHECK_JSON")
CACHE_DB_PATH = required_path_env("EVAL_CACHE_DB_PATH")
CACHE_EARLIER_PATH = required_path_env("EVAL_CACHE_DB_PATH_EARLIER")
SCOPUS_XLSX_PATH = required_path_env("EVAL_SCOPUS_XLSX_PATH")
OUTPUT_DIR = EVAL_OUTPUT_DIR / "output_sibling_issnl_scopus_figures"
stats_common.OUTPUT_DIR = OUTPUT_DIR
NUMBERS_DIR = stats_common.NUMBERS_DIR

# start time of the assessment run: a row carrying an ISSN that is only in the later cache must have been created at or after it
RUN_START = os.environ.get("EVAL_ASSESSMENT_RUN_START")
if not RUN_START:
    raise RuntimeError(
        "EVAL_ASSESSMENT_RUN_START is not set. Set it to the start time of the assessment run, as 'YYYY-MM-DD HH:MM:SS' "
        "in the format of the created_at column of the cache."
    )

# the tool's column-name mappings for the Scopus workbook (scopus_column_mappings)
SCOPUS_COLUMN_MAPPINGS = {
    "title": ["source title", "title"],
    "issn": ["issn"],
    "eissn": ["eissn", "e-issn"],
    "publisher": ["publisher"],
    "status": ["active or inactive"],
    "quality_flag": ["discontinued", "quality"],
    "source_type": ["source type"],
    "coverage": ["coverage"],
    "open_access": ["open access"],
}


def stored_number(stem: str, key: str) -> int:
    d = json.loads((NUMBERS_DIR / f"{stem}.numbers.json").read_text())[key]
    assert d["kind"] == "int", (stem, key)
    return d["value"]


def count_sibling_inferred(conn, strict_exclude: set, broad_drop: set) -> dict:
    # stats_corpus.py: files WHERE category='assessed'; sibling = no-log file whose submission has a file with a log
    frows = conn.execute("SELECT month, name, has_cited_sidecar FROM files WHERE category='assessed'").fetchall()
    sub_of = lambda m, n: (m, submission_of(n))  # noqa: E731
    subs_with_log = {sub_of(r["month"], r["name"]) for r in frows if r["has_cited_sidecar"]}
    sib_files = {(r["month"], r["name"]) for r in frows if not r["has_cited_sidecar"] and sub_of(r["month"], r["name"]) in subs_with_log}
    assert len(sib_files) == stored_number("stats_corpus", "corpus_nosidecar_sibling_files_n"), len(sib_files)

    rows = conn.execute(
        "SELECT month, arxiv_id, suffix, entry_key, final_assessment, n_predatory_list_hits, citation_status, citation_status_bbl_tex "
        "FROM entries WHERE citation_status_bbl_tex_changed = '1' AND citation_status = 'excluded_sibling_inferred_uncited' "
        "AND citation_status_bbl_tex = 'confirmed_cited'"
    ).fetchall()
    n_entries = len(rows)
    assert n_entries == stored_number("confirmed_cited_composition", "confpop_n_promoted_no_log_sibling"), n_entries
    files, subs = set(), set()
    n_strict = n_broad = 0
    for r in rows:
        f = (r["month"], f'{r["arxiv_id"]}-{r["suffix"]}')
        assert f in sib_files, "a promoted sibling-inferred entry sits in a file that is not a sibling file"
        files.add(f)
        subs.add((r["month"], r["arxiv_id"]))
        k = (r["month"], r["arxiv_id"], r["suffix"], r["entry_key"])
        n_strict += r["final_assessment"] == "predatory" and k not in strict_exclude
        n_broad += int(r["n_predatory_list_hits"] or 0) > 0 and k not in broad_drop

    # sibling files that hold no excluded entry, and those whose every entry is a promoted one
    per_file: dict = defaultdict(lambda: {"n": 0, "excluded": 0, "promoted": 0})
    for r in conn.execute("SELECT month, arxiv_id, suffix, citation_status_bbl_tex, citation_status_bbl_tex_changed FROM entries"):
        f = (r["month"], f'{r["arxiv_id"]}-{r["suffix"]}')
        if f in sib_files:
            p = per_file[f]
            p["n"] += 1
            p["excluded"] += r["citation_status_bbl_tex"] == "excluded_sibling_inferred_uncited"
            p["promoted"] += r["citation_status_bbl_tex_changed"] == "1"
    assert len(per_file) == len(sib_files), "a sibling file without any entry row"
    no_excluded = {f for f, p in per_file.items() if p["excluded"] == 0}
    all_promoted = {f for f, p in per_file.items() if p["promoted"] == p["n"]}
    assert no_excluded == all_promoted, "the sibling files without an excluded entry are not exactly the all-promoted files"
    return {
        "n_files_total": len(sib_files), "n_promoted_entries": n_entries, "n_promoted_files": len(files),
        "n_promoted_submissions": len(subs), "n_promoted_strict": n_strict, "n_promoted_broad": n_broad,
        "n_files_without_excluded_entry": len(no_excluded),
    }


def count_issnl_and_cache(strict_base: int, broad_base: int, keys: dict) -> dict:
    mism = {tuple(k) for k in keys["confirmed_misattribution_keys"]}
    survivors = {tuple(k) for k in keys["second_vote_survivor_keys"]}
    strict_exclude = {tuple(k) for k in keys["strict_exclude_keys"]}
    broad_drop = {tuple(k) for k in keys["broad_drop_keys"]}
    d = json.loads(ISSN_L_CHECK_JSON.read_text())
    detail = {tuple(r["key"]): r for r in d["per_citation_detail"]}
    if set(detail) != mism or d["n_confirmed_misattribution_citations"] != len(mism):
        raise SystemExit("the ISSN-L family check JSON is not on the basis of the audit keys file "
                         f"({len(detail)} vs {len(mism)} mismatches)")

    n_diff = n_null_only = 0
    for r in detail.values():
        types = {b["type"] for b in r["invalid_votes"]}
        assert types <= {"issn_mismatch", "null_cache_issn"}, types
        if "issn_mismatch" in types:
            n_diff += 1
        else:
            n_null_only += 1

    dist = d["pair_resolution_distribution"]
    assert sum(dist.values()) == d["n_distinct_issn_mismatch_pairs"]

    included_again = {tuple(k) for k in d["restored_citation_keys"]}
    assert included_again <= strict_exclude and not (included_again & survivors)

    # the cache question: references with no "cache ISSN differs" vote whose citing ISSN is on no row of the earlier cache
    later = sqlite3.connect(f"file:{CACHE_DB_PATH}?mode=ro", uri=True)
    earlier = sqlite3.connect(f"file:{CACHE_EARLIER_PATH}?mode=ro", uri=True)
    idx_later = build_any_row_has_issn_index(later)
    idx_earlier = build_any_row_has_issn_index(earlier)
    # created_at of the rows that carry an ISSN in the later cache
    created: dict = defaultdict(list)
    for issn, c in later.execute("SELECT issn, created_at FROM journals WHERE issn IS NOT NULL AND issn != ''"):
        created[normalize_issn(issn)].append(c)
    only_later: set = set()
    for k, r in detail.items():
        types = {b["type"] for b in r["invalid_votes"]}
        if "issn_mismatch" in types:
            continue
        ci = normalize_issn(r["citing_issn"])
        assert ci in idx_later, "a null-cache mismatch whose citing ISSN is on no row of the later cache"
        if ci not in idx_earlier:
            assert all((c or "") >= RUN_START for c in created[ci]), "a row carrying a later-cache-only ISSN was created before the run start"
            only_later.add(k)
    assert not (only_later & included_again)
    assert only_later <= strict_exclude and not (only_later & survivors)

    def counts_with(extra: set) -> tuple[int, int]:
        return strict_base + len(extra & strict_exclude), broad_base + len(extra & broad_drop)

    strict_with_included_again, broad_with_included_again = counts_with(included_again)
    strict_without_later_cache_only, broad_without_later_cache_only = counts_with(only_later)
    strict_with_both, broad_with_both = counts_with(included_again | only_later)
    assert (strict_with_both - strict_base) == len(included_again) + len(only_later)
    return {
        "n_mismatch_refs": len(detail), "n_refs_cache_issn_differs": n_diff, "n_refs_null_cache_only": n_null_only,
        "n_pairs": d["n_distinct_issn_mismatch_pairs"], "n_pairs_same_family": dist.get("same_issnl_family", 0),
        "n_pairs_distinct_family": dist.get("confirmed_distinct_family", 0),
        "n_pairs_unresolved": dist.get("unresolved_not_in_openalex", 0),
        "n_restored": len(included_again), "strict_with_restored": strict_with_included_again,
        "broad_with_restored": broad_with_included_again,
        "n_later_cache_only": len(only_later), "strict_without_later_cache_only": strict_without_later_cache_only,
        "broad_without_later_cache_only": broad_without_later_cache_only,
        "strict_with_both": strict_with_both, "broad_with_both": broad_with_both, "strict_max_movement": len(included_again | only_later),
    }


def count_scopus() -> dict:
    from openpyxl import load_workbook  # only this script needs openpyxl

    wb = load_workbook(filename=SCOPUS_XLSX_PATH, read_only=True, data_only=True)
    sheet = None
    for name in wb.sheetnames:  # the tool's scopus.py: first sheet whose name has "scopus" or "source"
        if "scopus" in name.lower() or "source" in name.lower():
            sheet = wb[name]
            break
    sheet = sheet or wb.active
    rows_iter = iter(sheet.rows)
    headers = [c.value for c in next(rows_iter)]
    col: dict = {}
    for i, h in enumerate(headers):  # the tool's scopus.py: column lookup
        if not h:
            continue
        hl = str(h).lower()
        for field, possible in SCOPUS_COLUMN_MAPPINGS.items():
            if field in col:
                continue
            if field == "quality_flag":
                if all(kw in hl for kw in possible):
                    col[field] = i
            else:
                for p in possible:
                    if p in hl:
                        col[field] = i
                        break
    assert {"title", "status", "quality_flag"} <= set(col), col
    cell = lambda row, f: row[col[f]].value if f in col else None  # noqa: E731
    n_rows = active = inactive = flagged = flagged_loaded = 0
    for row in rows_iter:  # the tool's scopus.py: row rules
        title = cell(row, "title")
        if not title or len(str(title).strip()) < 2:
            continue
        n_rows += 1
        status = cell(row, "status")
        is_active = (str(status).strip() if status else "").lower() == "active"
        qf = cell(row, "quality_flag")
        is_flagged = bool(qf and str(qf).strip())
        active += is_active
        inactive += not is_active
        flagged += is_flagged
        flagged_loaded += is_flagged and is_active

    cache = sqlite3.connect(f"file:{CACHE_DB_PATH}?mode=ro", uri=True)
    stored = cache.execute(
        "SELECT su.records_updated FROM source_updates su JOIN data_sources ds ON ds.id = su.source_id "
        "WHERE ds.name = 'scopus' AND su.update_type = 'full' AND su.status = 'success' ORDER BY su.id DESC LIMIT 1"
    ).fetchone()
    assert stored is not None and stored[0] == active, f"loaded active titles {active} != cache source_updates {stored}"
    return {"n_rows": n_rows, "n_active": active, "n_inactive": inactive, "n_quality_flagged_rows": flagged,
            "n_quality_flagged_loaded": flagged_loaded, "n_loaded_in_cache_record": stored[0]}


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    keys = json.loads(AUDIT_KEYS_JSON.read_text())
    strict_exclude = {tuple(k) for k in keys["strict_exclude_keys"]}
    broad_drop = {tuple(k) for k in keys["broad_drop_keys"]}
    conn = sqlite3.connect(f"file:{DB_PATH_BBL_TEX}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    strict_count_without_identity_check = conn.execute(
        "SELECT COUNT(*) FROM entries WHERE citation_status_bbl_tex = 'confirmed_cited' AND final_assessment = 'predatory'").fetchone()[0]
    broad_count_without_identity_check = conn.execute(
        "SELECT COUNT(*) FROM entries WHERE citation_status_bbl_tex = 'confirmed_cited' AND CAST(n_predatory_list_hits AS INTEGER) > 0").fetchone()[0]
    strict_base = strict_count_without_identity_check - len(strict_exclude)
    broad_base = broad_count_without_identity_check - len(broad_drop)
    # strict and broad counts after the identity check
    assert strict_base > 0 and broad_base > 0

    sibling = count_sibling_inferred(conn, strict_exclude, broad_drop)
    issnl = count_issnl_and_cache(strict_base, broad_base, keys)
    scopus = count_scopus()
    out = {"strict_base": strict_base, "broad_base": broad_base, "sibling_promoted": sibling, "issnl_and_cache": issnl, "scopus": scopus}
    stats_common.write_json("sibling_issnl_scopus_figures", out)
    print(json.dumps(out, indent=2))

    i = lambda v: {"kind": "int", "value": int(v)}  # noqa: E731
    numbers = {"sisf_strict_base_n": i(strict_base), "sisf_broad_base_n": i(broad_base)}
    numbers.update({
        "sisf_sib_n_files_total": i(sibling["n_files_total"]),
        "sisf_sib_n_promoted_entries": i(sibling["n_promoted_entries"]),
        "sisf_sib_n_promoted_files": i(sibling["n_promoted_files"]),
        "sisf_sib_n_promoted_submissions": i(sibling["n_promoted_submissions"]),
        "sisf_sib_n_promoted_strict": i(sibling["n_promoted_strict"]),
        "sisf_sib_n_promoted_broad": i(sibling["n_promoted_broad"]),
        "sisf_sib_n_files_without_excluded_entry": i(sibling["n_files_without_excluded_entry"]),
    })
    for k in ("n_mismatch_refs", "n_refs_cache_issn_differs", "n_refs_null_cache_only", "n_pairs", "n_pairs_same_family",
              "n_pairs_distinct_family", "n_pairs_unresolved", "n_restored", "strict_with_restored", "broad_with_restored",
              "n_later_cache_only", "strict_without_later_cache_only", "broad_without_later_cache_only",
              "strict_with_both", "broad_with_both", "strict_max_movement"):
        numbers[f"sisf_issnl_{k}" if not k.startswith(("strict_", "broad_")) else f"sisf_issnl_{k}_n"] = i(issnl[k])
    numbers["sisf_issnl_strict_max_movement_pct"] = {
        "kind": "ratio", "num": int(issnl["strict_max_movement"]), "den": int(strict_base), "decimals": 3}
    for k, v in scopus.items():
        numbers[f"sisf_scopus_{k}"] = i(v)
    stats_common.write_numbers("sibling_issnl_scopus_figures", numbers)


if __name__ == "__main__":
    main()
